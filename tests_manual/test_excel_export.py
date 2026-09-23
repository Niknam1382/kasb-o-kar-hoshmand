"""
تست: باگ #3 — خروجیِ اکسل باید یه فایلِ xlsx معتبر و استایل‌خورده باشه (راست‌به‌چپ،
هدرِ رنگی و بولد، بردر، ستون‌های هم‌عرض‌شده) نه یه دامپِ خامِ بدونِ فرمت.

اجرا: python3 tests_manual/test_excel_export.py
"""
from __future__ import annotations

import asyncio
import io
import sys

sys.path.insert(0, ".")
sys.path.insert(0, "tests_manual")

from openpyxl import load_workbook  # noqa: E402

from harness import reset_database  # noqa: E402
from shop_setup import seed_usable_shop  # noqa: E402

from app.database.models import Customer, OrderConsultation, OrderType  # noqa: E402
from app.database.session import session_scope  # noqa: E402
from app.services import excel_export_service  # noqa: E402

OWNER_TG_ID = 999501
SHOP_BOT_TG_ID = 999502


async def _seed_customers_and_orders(shop_bot_id: int) -> None:
    async with session_scope() as session:
        c1 = Customer(shop_bot_id=shop_bot_id, telegram_id=1001, first_name="نیلوفر", username="niloofar_shop")
        c2 = Customer(shop_bot_id=shop_bot_id, telegram_id=1002, first_name="کیوان", username=None)
        session.add_all([c1, c2])
        await session.flush()

        session.add_all(
            [
                OrderConsultation(
                    shop_bot_id=shop_bot_id, customer_id=c1.id, type=OrderType.ORDER, summary="خریدِ ۲ عدد کیف", estimated_value_toman=900_000
                ),
                OrderConsultation(
                    shop_bot_id=shop_bot_id, customer_id=c1.id, type=OrderType.CONSULTATION, summary="سوال درباره‌ی رنگ‌بندی", estimated_value_toman=None
                ),
                OrderConsultation(
                    shop_bot_id=shop_bot_id, customer_id=c2.id, type=OrderType.ORDER, summary="خریدِ کفش", estimated_value_toman=1_200_000
                ),
            ]
        )


async def test_customers_workbook_is_styled_and_correct() -> None:
    await reset_database()
    shop_bot_id = await seed_usable_shop(OWNER_TG_ID, SHOP_BOT_TG_ID)
    await _seed_customers_and_orders(shop_bot_id)

    async with session_scope() as session:
        raw_bytes = await excel_export_service.build_customers_workbook(session, shop_bot_id)

    assert isinstance(raw_bytes, bytes) and len(raw_bytes) > 0, "باید یه محتوای بایتیِ غیرخالی برگرده"
    assert raw_bytes[:2] == b"PK", "فایلِ xlsx باید با امضای ZIP (PK) شروع بشه"

    wb = load_workbook(io.BytesIO(raw_bytes))
    ws = wb.active

    # --- چکِ استایل (اصلِ باگ‌فیکس) ---
    assert ws.sheet_view.rightToLeft is True, "شیت باید راست‌به‌چپ باشه"
    assert ws.freeze_panes == "A2", "ردیفِ هدر باید فریز شده باشه"
    header_cell = ws["A1"]
    assert header_cell.font.bold is True, "فونتِ هدر باید بولد باشه"
    assert header_cell.fill.start_color.rgb in ("002F5496", "FF2F5496", "2F5496"), f"رنگِ پس‌زمینه‌ی هدر: {header_cell.fill.start_color.rgb}"
    assert header_cell.border.left.style == "thin", "سلول‌ها باید بردر داشته باشن"

    body_cell = ws["A2"]
    assert body_cell.alignment.horizontal == "right", "متنِ بدنه باید راست‌چین باشه"

    # حداقل یه ردیف باید رنگِ متناوب (zebra striping) داشته باشه
    row2_fill = ws["A2"].fill.start_color.rgb
    row3_fill = ws["A3"].fill.start_color.rgb
    assert row2_fill != row3_fill or (ws["A3"].fill.fill_type is None) != (ws["A2"].fill.fill_type is None), (
        "ردیف‌های زوج/فرد باید رنگ‌بندیِ متفاوت داشته باشن (zebra striping)"
    )

    # --- چکِ داده ---
    header_row = [cell.value for cell in ws[1]]
    assert header_row[0] == "نام"
    assert "تعداد سفارش" in header_row and "تعداد مشاوره" in header_row and "مجموع ارزش تخمینی (تومان)" in header_row

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2, f"باید دقیقاً دو ردیفِ مشتری باشه، نه {len(rows)}"

    niloofar_row = next(r for r in rows if r[0] == "نیلوفر")
    assert niloofar_row[1] == "@niloofar_shop"
    assert niloofar_row[5] == 1, f"نیلوفر باید یه سفارش داشته باشه: {niloofar_row[5]}"
    assert niloofar_row[6] == 1, f"نیلوفر باید یه مشاوره داشته باشه: {niloofar_row[6]}"
    assert niloofar_row[7] == 900_000, f"مجموع ارزش: {niloofar_row[7]}"

    kayvan_row = next(r for r in rows if r[0] == "کیوان")
    assert kayvan_row[1] in ("", None), "بدونِ یوزرنیم باید سلولِ خالی باشه (openpyxl رشته‌ی خالی رو در ذخیره/بارگذاری None می‌کنه)"
    assert kayvan_row[5] == 1
    assert kayvan_row[7] == 1_200_000

    print("✅ test_customers_workbook_is_styled_and_correct PASSED")


async def test_orders_workbook_is_valid() -> None:
    await reset_database()
    shop_bot_id = await seed_usable_shop(OWNER_TG_ID + 1, SHOP_BOT_TG_ID + 1)
    await _seed_customers_and_orders(shop_bot_id)

    async with session_scope() as session:
        raw_bytes = await excel_export_service.build_orders_workbook(session, shop_bot_id)

    wb = load_workbook(io.BytesIO(raw_bytes))
    ws = wb.active
    assert ws.sheet_view.rightToLeft is True
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 3, f"باید سه ردیفِ سفارش/مشاوره باشه، نه {len(rows)}"

    print("✅ test_orders_workbook_is_valid PASSED")


async def main() -> None:
    await test_customers_workbook_is_styled_and_correct()
    await test_orders_workbook_is_valid()


if __name__ == "__main__":
    asyncio.run(main())
