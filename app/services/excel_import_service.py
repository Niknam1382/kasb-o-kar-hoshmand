from __future__ import annotations

import io
import re
from dataclasses import dataclass

from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.database.models import Product
from app.services import product_service
from app.services.excel_export_service import _style_sheet
from app.utils.validators import parse_price_toman

# حداکثرِ تعدادِ ردیفِ دادهٔ قابل‌قبول در یه فایل — محافظت در برابرِ فایل‌هایِ
# بیش‌ازحد بزرگ (هم برایِ عملکرد، هم چون هیچ فروشگاهِ کوچیکی واقعاً بیشتر از این
# به یه بارگذاریِ تکی نیاز نداره؛ اگه لازم شد، کاربر فایل رو چندتکه می‌کنه).
MAX_IMPORT_ROWS = 500

TEMPLATE_HEADERS = ["نام محصول", "توضیحات", "قیمت (تومان)", "موجودی (خالی = نامحدود)"]

_STOCK_CLEAN_RE = re.compile(r"[,\u060C\s]")


class InvalidExcelFileError(Exception):
    """فایل اصلاً به‌عنوانِ xlsx قابلِ بازشدن نیست (خراب/فرمتِ اشتباه)."""


@dataclass
class ParseResult:
    valid_rows: list[dict]
    errors: list[str]
    too_many_rows: bool = False


def build_template_workbook() -> bytes:
    """یه فایلِ اکسلِ خالی (فقط هدر، استایل‌خورده) برای پرکردن توسطِ فروشگاه‌دار می‌سازه."""
    wb = Workbook()
    ws = wb.active
    ws.title = "محصولات"
    ws.append(TEMPLATE_HEADERS)
    _style_sheet(ws)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _cell_str(value: object) -> str:
    """مقدارِ یه سلولِ اکسل (هرنوعی: str/int/float/None) رو به رشته‌ی trimشده تبدیل
    می‌کنه. اعدادِ اعشاریِ صحیح (مثلِ 150000.0 که openpyxl گاهی برمی‌گردونه) بدونِ
    ممیز نمایش داده می‌شن."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_stock(raw: str) -> int | None:
    """raw رو به عددِ صحیحِ غیرمنفی تبدیل می‌کنه؛ اگه نامعتبر بود None برمی‌گردونه.
    برخلافِ قیمت، صفر برایِ موجودی مقدارِ معتبریه (یعنی «ناموجود ولی پیگیری‌شونده»)."""
    cleaned = _STOCK_CLEAN_RE.sub("", raw)
    if not cleaned.isdigit():
        return None
    return int(cleaned)


def _parse_row(row: tuple, row_number: int) -> tuple[dict | None, str | None]:
    """یه ردیفِ خامِ اکسل رو اعتبارسنجی می‌کنه.

    برمی‌گردونه:
    - (دیکشنریِ محصولِ معتبر، None) اگه ردیف درست بود
    - (None, پیامِ خطایِ فارسی) اگه ردیف مشکل داشت
    - (None, None) اگه ردیف کاملاً خالی بود (نادیده گرفته می‌شه، نه خطا — معمولاً
      ردیف‌های اضافیِ خالیِ ته‌شیت هستن)
    """
    cells = list(row) + [None, None, None, None]
    name_raw = _cell_str(cells[0])
    desc_raw = _cell_str(cells[1])
    price_raw = _cell_str(cells[2])
    stock_raw = _cell_str(cells[3])

    if not (name_raw or desc_raw or price_raw or stock_raw):
        return None, None

    if not name_raw:
        return None, f"ردیف {row_number}: نامِ محصول خالیه."
    if len(name_raw) > 255:
        return None, f"ردیف {row_number}: نامِ محصول بیش از ۲۵۵ کاراکتره."

    if not price_raw:
        return None, f"ردیف {row_number} («{name_raw}»): قیمت خالیه."
    price = parse_price_toman(price_raw)
    if price is None:
        return None, f"ردیف {row_number} («{name_raw}»): قیمت نامعتبره — فقط عددِ مثبت بنویس."

    stock: int | None = None
    if stock_raw:
        stock = _parse_stock(stock_raw)
        if stock is None:
            return None, f"ردیف {row_number} («{name_raw}»): موجودی نامعتبره — باید عددِ صحیحِ صفر یا بزرگ‌تر باشه."

    return {
        "name": name_raw,
        "description": desc_raw or None,
        "price_toman": price,
        "stock_quantity": stock,
    }, None


def parse_workbook(content: bytes) -> ParseResult:
    """محتوایِ بایتیِ یه فایلِ xlsx آپلودشده رو می‌خونه و ردیف‌به‌ردیف اعتبارسنجی می‌کنه.

    فرضِ ستون‌ها دقیقاً مثلِ TEMPLATE_HEADERS هست (نام/توضیحات/قیمت/موجودی)؛ ردیفِ
    اول (هدر) همیشه نادیده گرفته می‌شه. عکس از این راه پشتیبانی نمی‌شه.
    """
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        ws = wb.active
        if ws is None:
            raise InvalidExcelFileError("این فایل هیچ شیتی نداره.")
    except InvalidExcelFileError:
        raise
    except Exception as exc:
        raise InvalidExcelFileError("فایل اکسل قابلِ خوندن نیست.") from exc

    valid_rows: list[dict] = []
    errors: list[str] = []
    processed = 0

    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        parsed, error = _parse_row(row, row_number)
        if parsed is None and error is None:
            continue

        processed += 1
        if processed > MAX_IMPORT_ROWS:
            return ParseResult(valid_rows=[], errors=[], too_many_rows=True)

        if error is not None:
            errors.append(error)
        else:
            valid_rows.append(parsed)

    return ParseResult(valid_rows=valid_rows, errors=errors)


async def bulk_create(session: AsyncSession, shop_bot_id: int, rows: list[dict]) -> list[Product]:
    """همه‌ی ردیف‌هایِ ازقبل‌اعتبارسنجی‌شده رو به‌عنوانِ محصولِ جدید ذخیره می‌کنه.
    چون همه‌چیز روی یه session ی مشترک (تزریقِ DbSessionMiddleware) اجرا می‌شه که
    فقط یه‌بار در انتهایِ همون آپدیت commit می‌شه، این عملیات به‌صورتِ طبیعی اتمیکه:
    یا همه‌ی محصولات ذخیره می‌شن، یا (اگه خطایی پیش بیاد) هیچ‌کدوم."""
    created: list[Product] = []
    for row in rows:
        product = await product_service.create_product(
            session,
            shop_bot_id,
            row["name"],
            row.get("description"),
            row["price_toman"],
            None,
            row.get("stock_quantity"),
        )
        created.append(product)
    return created
