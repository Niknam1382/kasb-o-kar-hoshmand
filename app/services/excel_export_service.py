from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database.models import Customer, OrderConsultation, OrderType, Product

_HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(fill_type="solid", start_color="2F5496", end_color="2F5496")
_BODY_FONT = Font(name="Arial")
_RTL_ALIGNMENT = Alignment(horizontal="right", readingOrder=2, vertical="center", wrap_text=True)
_HEADER_ALIGNMENT = Alignment(horizontal="center", readingOrder=2, vertical="center", wrap_text=True)
_THIN_SIDE = Side(style="thin", color="BFBFBF")
_CELL_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)
_ALT_ROW_FILL = PatternFill(fill_type="solid", start_color="F2F2F2", end_color="F2F2F2")


def _style_sheet(ws: Worksheet) -> None:
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"

    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _CELL_BORDER
    ws.row_dimensions[1].height = 24

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for cell in row:
            cell.font = _BODY_FONT
            cell.alignment = _RTL_ALIGNMENT
            cell.border = _CELL_BORDER
            if row_idx % 2 == 0:
                cell.fill = _ALT_ROW_FILL

    for col in ws.columns:
        max_len = max((len(str(cell.value)) for cell in col if cell.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 4, 12), 45)


def _fmt(dt) -> str:
    return dt.strftime("%Y-%m-%d %H:%M") if dt else ""


async def build_customers_workbook(session: AsyncSession, shop_bot_id: int) -> bytes:
    result = await session.execute(select(Customer).where(Customer.shop_bot_id == shop_bot_id).order_by(Customer.id))
    customers = result.scalars().all()

    # تعداد سفارش/مشاوره و مجموع ارزش تخمینی به‌ازای هر مشتری، با یک کوئری گروه‌بندی‌شده
    stats_result = await session.execute(
        select(
            OrderConsultation.customer_id,
            OrderConsultation.type,
            func.count(),
            func.coalesce(func.sum(OrderConsultation.estimated_value_toman), 0),
        )
        .where(OrderConsultation.shop_bot_id == shop_bot_id)
        .group_by(OrderConsultation.customer_id, OrderConsultation.type)
    )
    per_customer: dict[int, dict[str, int]] = {}
    for customer_id, order_type, count, total_value in stats_result.all():
        key = order_type.value if hasattr(order_type, "value") else order_type
        entry = per_customer.setdefault(customer_id, {"order": 0, "consultation": 0, "value": 0})
        entry[key] = count
        entry["value"] += int(total_value or 0)

    wb = Workbook()
    ws = wb.active
    ws.title = "مشتریان"
    ws.append(
        [
            "نام",
            "یوزرنیم",
            "آیدی تلگرام",
            "اولین پیام",
            "آخرین پیام",
            "تعداد سفارش",
            "تعداد مشاوره",
            "مجموع ارزش تخمینی (تومان)",
        ]
    )

    for c in customers:
        stats = per_customer.get(c.id, {"order": 0, "consultation": 0, "value": 0})
        ws.append(
            [
                c.first_name or "",
                f"@{c.username}" if c.username else "",
                c.telegram_id,
                _fmt(c.first_seen_at),
                _fmt(c.last_message_at),
                stats["order"],
                stats["consultation"],
                stats["value"],
            ]
        )

    _style_sheet(ws)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def build_orders_workbook(session: AsyncSession, shop_bot_id: int) -> bytes:
    result = await session.execute(
        select(OrderConsultation, Customer, Product)
        .join(Customer, OrderConsultation.customer_id == Customer.id)
        .outerjoin(Product, OrderConsultation.product_id == Product.id)
        .where(OrderConsultation.shop_bot_id == shop_bot_id)
        .order_by(OrderConsultation.id.desc())
    )
    rows = result.all()

    wb = Workbook()
    ws = wb.active
    ws.title = "سفارش‌ها و مشاوره‌ها"
    ws.append(
        [
            "تاریخ",
            "نوع",
            "مشتری",
            "آیدی تلگرام مشتری",
            "محصول",
            "تعداد",
            "خلاصه",
            "ارزش تخمینی (تومان)",
            "وضعیت تایید",
        ]
    )

    type_labels = {"order": "سفارش", "consultation": "مشاوره"}
    for order, customer, product in rows:
        order_type = order.type.value if hasattr(order.type, "value") else order.type
        if order_type == OrderType.ORDER.value or product is not None:
            confirmed_label = "✅ تایید و کسر شده" if order.confirmed else "⏳ در انتظار تایید"
        else:
            confirmed_label = "—"

        ws.append(
            [
                _fmt(order.created_at),
                type_labels.get(order_type, order_type),
                customer.first_name or "",
                customer.telegram_id,
                product.name if product is not None else "",
                order.quantity or "",
                order.summary,
                order.estimated_value_toman or "",
                confirmed_label,
            ]
        )

    _style_sheet(ws)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
